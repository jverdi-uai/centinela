from pathlib import Path
from statistics import median

from openpyxl import load_workbook

from centinela.agents.documents import DocumentAgent


def main() -> None:
    workbook = load_workbook("data/synthetic/datos_prueba_centinela.xlsx", read_only=True, data_only=True)
    rows = list(workbook["documentos_prueba"].iter_rows(min_row=2, values_only=True))
    agent = DocumentAgent()
    outcomes = []
    for filename, document_type, label, *_ in rows:
        path = Path("data/samples") / filename
        decision = agent.evaluate({"filename": filename, "content": path.read_bytes(), "document_type": document_type, "expected_fields": None, "reference_id": None})
        outcomes.append((label, decision))
    correct = sum((label == "autentico" and d.verdict == "autentico") or (label == "falso" and d.verdict in {"sospechoso", "falso"}) for label, d in outcomes)
    fraud = [(label, d) for label, d in outcomes if label == "falso"]
    recall = sum(d.verdict in {"sospechoso", "falso"} for _, d in fraud) / len(fraud)
    failures: dict[str, int] = {}
    for _, decision in outcomes:
        for check in decision.checks:
            if check.status == "fail":
                failures[check.name] = failures.get(check.name, 0) + 1
    latencies = sorted(d.latency_ms for _, d in outcomes)
    cost = sum(d.cost_usd for _, d in outcomes)
    report = f"""# Evaluación de documentos

Modo: `{agent.settings.mock_mode and 'mock' or 'OpenAI'}` · Casos: {len(outcomes)}

| Métrica | Resultado |
|---|---:|
| Exactitud operativa | {correct / len(outcomes):.3f} |
| Recall de documentos alterados | {recall:.3f} |
| Latencia p50 | {median(latencies):.2f} ms |
| Latencia p95 | {latencies[round((len(latencies)-1)*.95)]:.2f} ms |
| Costo por documento | USD {cost / len(outcomes):.8f} |

Checks fallidos más frecuentes: {failures or 'ninguno'}.

La comparación entre modelos queda disponible al ejecutar dos veces con `OPENAI_VISION_MODEL`; no se consume API en la evaluación mock.
"""
    Path("reports/documents_eval.md").write_text(report, encoding="utf-8")
    print(report)


if __name__ == "__main__":
    main()

