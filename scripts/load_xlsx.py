"""Carga el Excel entregado y reconstruye TransactionInput desde su formato aplanado."""
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from centinela.agents.transactions.schemas import TransactionInput


def excel_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().upper() in {"TRUE", "=TRUE()", "VERDADERO", "=VERDADERO()"}
    return bool(value)


def load_transactions(path: str | Path) -> list[tuple[TransactionInput, str, str | None]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook["transacciones"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    output: list[tuple[TransactionInput, str, str | None]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values, strict=True))
        transaction = TransactionInput.model_validate({
            "transaction_id": row["transaction_id"], "timestamp": row["timestamp"],
            "amount": row["amount"], "currency": row["currency"], "channel": row["channel"], "type": row["type"],
            "origin_account": {"id": row["origin_account_id"], "age_days": row["origin_age_days"], "avg_monthly_amount": row["origin_avg_monthly_amount"], "country": row["origin_country"]},
            "destination_account": {"id": row["destination_account_id"], "bank": row["destination_bank"], "is_new_beneficiary": excel_bool(row["is_new_beneficiary"]), "country": row["destination_country"]},
            "device": {"id": row["device_id"], "is_new_device": excel_bool(row["is_new_device"]), "os": row["os"], "ip_country": row["ip_country"], "vpn": excel_bool(row["vpn"])},
            "geo": {"lat": -33.45, "lon": -70.66, "distance_from_home_km": row["distance_from_home_km"]},
            "behavior": {"tx_last_hour": row["tx_last_hour"], "tx_last_24h": row["tx_last_24h"], "failed_logins_24h": row["failed_logins_24h"], "session_seconds": row["session_seconds"]},
            "customer_profile": {"segment": row["segment"], "risk_tier": row["risk_tier"]},
        })
        output.append((transaction, row["etiqueta"], row["patron"]))
    return output


if __name__ == "__main__":
    source = Path("data/synthetic/datos_prueba_centinela.xlsx")
    rows = load_transactions(source)
    print(f"Cargadas {len(rows)} transacciones desde {source}")

